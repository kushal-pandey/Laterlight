
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as WordDocument
from fastapi import HTTPException
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def extract_document_text(filename: str, raw_bytes: bytes) -> str:
    """Extract readable text from a supported document."""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Upload TXT, PDF, or DOCX.",
        )

    try:
        if extension == ".txt":
            text = raw_bytes.decode("utf-8-sig")

        elif extension == ".pdf":
            reader = PdfReader(BytesIO(raw_bytes))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        else:
            document = WordDocument(BytesIO(raw_bytes))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Include text from tables as well.
            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

            text = "\n".join(parts)

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="The uploaded document could not be read. "
            "Check that it is a valid, uncorrupted file.",
        ) from None

    if not text.strip():
        if extension == ".pdf":
            raise HTTPException(
                status_code=422,
                detail="No selectable text found in this PDF. "
                "Scanned PDFs require OCR, which is not supported yet.",
            )

        raise HTTPException(
            status_code=422,
            detail="No readable text found in the uploaded document.",
        )

    return text
